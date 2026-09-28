import re, sys, copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

SRC, OUT = sys.argv[1], sys.argv[2]
ICON = "icons/{}_{}.png"

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
SAFF = RGBColor(0xE0, 0x6C, 0x14)
GREEN = RGBColor(0x2E, 0x8B, 0x3E)
TEAL = RGBColor(0x0F, 0x76, 0x6E)
RED = RGBColor(0xB4, 0x3C, 0x2E)
GREY = RGBColor(0x26, 0x26, 0x26)
MUTED = RGBColor(0x55, 0x60, 0x6E)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
TINT = RGBColor(0xF1, 0xF5, 0xFA)
TINT_S = RGBColor(0xFD, 0xF1, 0xE7)
TINT_G = RGBColor(0xEA, 0xF5, 0xEC)
LINE = RGBColor(0xC9, 0xD3, 0xDF)
FONT = "Arial"

prs = Presentation(SRC)
S = prs.slides


def rm(shape):
    el = shape._element
    el.getparent().remove(el)


def find(slide, name):
    return [s for s in slide.shapes if s.name == name][0]


def box(sl, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06, lw=0.75):
    s = sl.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line; s.line.width = Pt(lw)
    s.shadow.inherit = False
    s.text_frame.text = ""
    return s


def fill_tf(tf, paras, size=11, color=GREY, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, margin=0.06, bold=False):
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    tf.vertical_anchor = anchor
    first = True
    for p in paras:
        if isinstance(p, str):
            p = {"t": p}
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        para.alignment = p.get("align", align)
        para.space_after = Pt(p.get("after", 3))
        sz = p.get("size", size)
        col = p.get("color", color)
        for i, seg in enumerate(re.split(r"(\*\*[^*]+\*\*)", p["t"])):
            if not seg:
                continue
            r = para.add_run()
            isb = seg.startswith("**")
            r.text = seg[2:-2] if isb else seg
            f = r.font
            f.name = FONT; f.size = Pt(sz); f.color.rgb = p.get("bcolor", col) if isb else col
            f.bold = True if isb else p.get("bold", bold)
            f.italic = p.get("italic", False)
        if p.get("bullet"):
            pPr = para._p.get_or_add_pPr()
            ind = Inches(p.get("indent", 0.16))
            pPr.set("marL", str(ind)); pPr.set("indent", str(-ind))
            bc = etree.SubElement(pPr, qn("a:buClr"))
            etree.SubElement(bc, qn("a:srgbClr")).set("val", p.get("bucol", "E06C14"))
            etree.SubElement(pPr, qn("a:buFont")).set("typeface", "Arial")
            etree.SubElement(pPr, qn("a:buChar")).set("char", p.get("char", "•"))


def text(sl, x, y, w, h, paras, **kw):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    fill_tf(tb.text_frame, paras, **kw)
    return tb


def label(sl, x, y, w, t, size=14, color=NAVY):
    """Official template pointer, kept verbatim as a section label."""
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(0.36))
    fill_tf(tb.text_frame, [{"t": "❖ ", "size": size}], color=color, margin=0)
    p = tb.text_frame.paragraphs[0]
    r = p.add_run(); r.text = t
    r.font.name = FONT; r.font.size = Pt(size); r.font.bold = True; r.font.underline = True; r.font.color.rgb = color
    p.runs[0].font.color.rgb = color; p.runs[0].font.bold = True
    return tb


def icon(sl, x, y, d, name, fill=NAVY):
    c = sl.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    c.fill.solid(); c.fill.fore_color.rgb = fill; c.line.fill.background(); c.shadow.inherit = False
    pad = d * 0.24
    sl.shapes.add_picture(ICON.format(name, "w"), Inches(x + pad), Inches(y + pad), Inches(d - 2 * pad), Inches(d - 2 * pad))


def card(sl, x, y, w, h, head, body, hcol=NAVY, fill=TINT, size=10.5, iconname=None):
    box(sl, x, y, w, h, fill=fill, line=LINE)
    hd = box(sl, x, y, w, 0.44, fill=hcol, radius=0.18)
    fill_tf(hd.text_frame, [{"t": head, "bold": True, "color": WHITE, "size": 12}],
            anchor=MSO_ANCHOR.MIDDLE, margin=0.12 if not iconname else 0.5)
    if iconname:
        sl.shapes.add_picture(ICON.format(iconname, "w"), Inches(x + 0.14), Inches(y + 0.08), Inches(0.28), Inches(0.28))
    text(sl, x + 0.08, y + 0.52, w - 0.16, h - 0.58, body, size=size)


def arrow(sl, x, y, w=0.26, h=0.3, col=MUTED):
    a = sl.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    a.fill.solid(); a.fill.fore_color.rgb = col; a.line.fill.background(); a.shadow.inherit = False


def B(t, **kw):
    d = {"t": t, "bullet": True, "after": kw.pop("after", 4)}
    d.update(kw)
    return d


TEAM_NAME, TEAM_ID = "TribeCoders", "144977"


def set_team(sl):
    for s in sl.shapes:
        if s.name.startswith("Oval") and s.has_text_frame and "Team" in s.text_frame.text:
            tf = s.text_frame
            tf.word_wrap = True
            tf.margin_left = tf.margin_right = Inches(0.02)
            tf.margin_top = tf.margin_bottom = Inches(0.02)
            p0 = tf.paragraphs[0]
            for extra in p0.runs[1:]:
                extra._r.getparent().remove(extra._r)
            r = p0.runs[0]
            r.text = TEAM_NAME
            r.font.size = Pt(12); r.font.bold = True
            p0.alignment = PP_ALIGN.CENTER
            p1 = tf.add_paragraph()
            p1.alignment = PP_ALIGN.CENTER
            r1 = p1.add_run()
            r1.text = f"ID {TEAM_ID}"
            r1.font.size = Pt(10)


# ---------------------------------------------------------------- remove instructions slide (7)
sldIdLst = prs.slides._sldIdLst
last = list(sldIdLst)[6]
prs.part.drop_rel(last.get(qn("r:id")))
sldIdLst.remove(last)

# ---------------------------------------------------------------- SLIDE 1: title page
s1 = S[0]
sub = find(s1, "Subtitle 3")
pars = sub.text_frame.paragraphs
for r in pars[-1].runs:
    if r.text.strip():
        r.text = "SAHARA"
p2 = copy.deepcopy(pars[-1]._p)
sub.text_frame._txBody.append(p2)
np_ = sub.text_frame.paragraphs[-1]
for r in np_.runs[1:]:
    r._r.getparent().remove(r._r)
np_.runs[0].text = "Predictive Personnel Stress & Welfare Intelligence"
np_.runs[0].font.size = Pt(18); np_.runs[0].font.bold = False; np_.runs[0].font.color.rgb = TEAL

tb = find(s1, "TextBox 9")
tb.top = Inches(2.55)
vals = {
    "Problem Statement ID": " SIH26186 (Ministry of Home Affairs)",
    "Problem Statement Title": " AI-Based Predictive Personnel Stress and Welfare Monitoring System for Uniformed Forces",
    "Theme": " MedTech / BioTech / HealthTech",
    "PS Category": None,
    "Team ID": " 144977",
    "Team Name": " TribeCoders",
}
for para in list(tb.text_frame.paragraphs):
    if not para.runs or not "".join(r.text for r in para.runs).strip():
        para._p.getparent().remove(para._p)
        continue
    para.alignment = PP_ALIGN.LEFT
    para.line_spacing = 1.0
    para.space_before = Pt(0)
    lab = para.runs[0]
    lab.font.size = Pt(18)
    para.space_after = Pt(10)
    key = next((k for k in vals if lab.text.startswith(k)), None)
    if key == "PS Category":
        lab.text = "PS Category-"
        v = " Software"
    elif key == "Team Name":
        lab.text = "Team Name (Registered on portal)-"
        v = vals[key]
    else:
        v = vals.get(key)
    if v:
        r = para.add_run(); r.text = v
        r.font.name = "Arial"; r.font.size = Pt(15 if key == "Problem Statement Title" else 16)
        r.font.bold = False; r.font.color.rgb = NAVY
s1.notes_slide.notes_text_frame.text = (
    "Opening line: SAHARA turns welfare from reactive to preventive. It predicts stress, burnout and emotional fatigue risk early, "
    "routes explainable alerts only to welfare staff, and proves recovery, fully privacy-first and on-premises.")

# ---------------------------------------------------------------- SLIDE 2: idea
s2 = S[1]
set_team(s2)
rm(find(s2, "TextBox 8"))
t = find(s2, "Title 1")
runs_ = [r for p in t.text_frame.paragraphs for r in p.runs]
runs_[0].text = ""
runs_[1].text = "SAHARA: Privacy-First Predictive Welfare Intelligence"
runs_[1].font.size = Pt(24)

label(s2, 0.4, 1.2, 12.5, "Proposed Solution (Describe your Idea/Solution/Prototype)", size=16)
text(s2, 0.4, 1.56, 12.5, 0.46, [{"t": "**730** CAPF, NSG & AR personnel died by suicide and **55,555** left service in 2020-24 (MHA, Rajya Sabha Q.1036). "
     "**SAHARA** predicts **stress, burnout, emotional fatigue** and **welfare concerns** early from HRMS + voluntary app data, and turns every alert into "
     "human-led support. **Welfare, never surveillance or discipline.**", "size": 10.5, "bcolor": NAVY}])

# Left: 6-step loop + the three official pointers as compact rows. Right: live prototype screenshot.
CAP_H = 0.58                                   # caption height under the screenshot
SH = 6.85 - CAP_H - 0.1 - 2.06                 # screenshot height that leaves room for the caption
SW = SH * 2400 / 1800
SX = 12.93 - SW
LW = SX - 0.18 - 0.4
steps = ["SENSE", "PROTECT", "PREDICT", "ALERT", "ACT", "RECOVER"]
cols = [NAVY, RGBColor(0x2A, 0x55, 0x7F), TEAL, SAFF, GREEN, RGBColor(0x1E, 0x6B, 0x2C)]
cw = (LW + 5 * 0.05) / 6
for i, a in enumerate(steps):
    shp = s2.shapes.add_shape(MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON, Inches(0.4 + i * (cw - 0.05)), Inches(2.06),
                              Inches(cw), Inches(0.46))
    shp.fill.solid(); shp.fill.fore_color.rgb = cols[i]; shp.line.fill.background(); shp.shadow.inherit = False
    fill_tf(shp.text_frame, [{"t": a, "bold": True, "size": 9, "color": WHITE, "align": PP_ALIGN.CENTER, "after": 0}],
            anchor=MSO_ANCHOR.MIDDLE, margin=0)

rows = [
    ("Detailed explanation of the proposed solution", NAVY, TINT, "FaLightbulb", [
        "**Mobile app:** consent per data type, 10-sec check-in, WHO-5, reaction test, camera heart-rate (on-device); Hindi, offline",
        "**HRMS analytics:** 26 features from leave, deployment, duty and night shifts, transfers, training, workload",
        "**AI engine:** personal baseline + anomaly + 4 calibrated XGBoost risk models + SHAP",
        "**Action:** tiered alerts, counterfactual advice, OR-Tools roster optimizer, recovery tracking"]),
    ("How it addresses the problem", TEAL, RGBColor(0xEC, 0xF6, 0xF5), "FaCheckCircle", [
        "**Late detection** → flagged early against each person's own baseline",
        "**Stigma & fear** → voluntary, pseudonymous; commanders see aggregates only",
        "**Alerts without action** → ranked interventions + fair roster + follow-up",
        "**Crisis moments** → Safe-Route to counsellor + Tele-MANAS 14416"]),
    ("Innovation and uniqueness of the solution", SAFF, TINT_S, "FaBrain", [
        "**Adaptive evidence:** one optional check lifts confidence 40% → 97% (live prototype)",
        "**'What would help' AI:** recovery window + leave → High to Moderate",
        "**My Data Mirror** + differential-privacy commander view",
        "**Purpose lock:** welfare data blocked from disciplinary use (403, audited)"]),
]
ry, rh, hw = 2.64, 1.36, 1.5
for head, col, tint, ic, pts in rows:
    box(s2, 0.4, ry, LW, rh, fill=tint, line=LINE)
    hd = box(s2, 0.4, ry, hw, rh, fill=col, radius=0.08)
    fill_tf(hd.text_frame, [{"t": head, "bold": True, "size": 10.5, "color": WHITE, "after": 0}],
            anchor=MSO_ANCHOR.MIDDLE, margin=0.1)
    s2.shapes.add_picture(ICON.format(ic, "w"), Inches(0.52), Inches(ry + 0.1), Inches(0.24), Inches(0.24))
    colw = (LW - hw - 0.2) / 2
    for c in range(2):
        text(s2, 0.4 + hw + 0.08 + c * (colw + 0.04), ry + 0.08, colw, rh - 0.12,
             [B(t, size=9.5, after=5, bucol="%02X%02X%02X" % (col[0], col[1], col[2])) for t in pts[2 * c: 2 * c + 2]], margin=0.03)
    ry += rh + 0.06

frame = box(s2, SX - 0.04, 2.02, SW + 0.08, SH + 0.08, fill=WHITE, line=RGBColor(0x9A, 0xA8, 0xBA), lw=1)
s2.shapes.add_picture("welfare_crop.png", Inches(SX), Inches(2.06), Inches(SW), Inches(SH))
cap = box(s2, SX - 0.04, 6.85 - CAP_H, SW + 0.08, CAP_H, fill=NAVY, radius=0.12)
fill_tf(cap.text_frame, [{"t": "Working prototype · Welfare Officer Console", "bold": True, "size": 10, "color": WHITE, "after": 1},
                         {"t": "Live screen from our running app (React + FastAPI + XGBoost), synthetic data. Case P-104: 4 risk scores, "
                               "evidence agreement, SHAP factors, counterfactual advice, human decision.", "size": 8,
                          "color": RGBColor(0xD9, 0xE3, 0xEF), "after": 0}], anchor=MSO_ANCHOR.MIDDLE, margin=0.12)
s2.notes_slide.notes_text_frame.text = (
    "Walk the 6-step loop, then point at the live screenshot: P-104's confidence rose from 40% to 97% after one optional "
    "check-in and readiness test. Show SHAP factors, the counterfactual (recovery window + leave), and that a human decides. "
    "Commanders never see this screen; they get differential-privacy aggregates only.")

# ---------------------------------------------------------------- SLIDE 3: technical approach
s3 = S[2]
set_team(s3)
rm(find(s3, "TextBox 8"))
label(s3, 0.4, 1.2, 4.2, "Technologies to be used", size=14)
tech = [
    ("FaMobileAlt", "Personnel app", "React 19 PWA, Hindi/English, on-device reaction test; Android via Capacitor"),
    ("FaLaptopCode", "Dashboards", "React 19, Vite, Tailwind CSS, Recharts"),
    ("FaServer", "Backend API", "Python FastAPI, Pydantic, JWT (OIDC-ready), OpenAPI"),
    ("FaDatabase", "Data", "PostgreSQL (production) · SQLite (prototype) via SQLAlchemy"),
    ("FaBrain", "AI / ML", "XGBoost x4 + isotonic calibration, SHAP, Isolation Forest, ruptures"),
    ("FaRoute", "Optimization", "Google OR-Tools CP-SAT fair-roster solver"),
    ("FaLock", "Security", "Role + unit + purpose checks, SHA-256 hash-chained audit, differential privacy"),
    ("FaCogs", "Deployment", "Force data centre or MeghRaj (GoI cloud); open-source, CPU-only, no public LLM"),
]
y = 1.62
for ic, h, d in tech:
    box(s3, 0.4, y, 4.2, 0.6, fill=TINT, line=LINE, radius=0.12)
    icon(s3, 0.48, y + 0.08, 0.44, ic, fill=NAVY if ic not in ("FaBrain", "FaRoute") else SAFF)
    text(s3, 1.0, y + 0.02, 3.55, 0.56, [{"t": h, "bold": True, "size": 10.5, "color": NAVY, "after": 0},
                                         {"t": d, "size": 9.5, "after": 0}], anchor=MSO_ANCHOR.MIDDLE, margin=0.02)
    y += 0.655

X0 = 4.85
label(s3, X0, 1.2, 8.1, "Methodology and process for implementation (Flow Charts/Images/ working prototype)", size=13)
colw, gap = 1.42, 0.27
xs = [X0 + i * (colw + gap) for i in range(5)]
top, bot = 1.66, 4.62
heads = ["DATA SOURCES", "PRIVACY", "AI ENGINE", "ACTION", "USERS"]
hc = [MUTED, NAVY, SAFF, TEAL, GREEN]
for i, hdr in enumerate(heads):
    text(s3, xs[i], top, colw, 0.26, [{"t": hdr, "bold": True, "size": 9.5, "color": hc[i], "align": PP_ALIGN.CENTER}], margin=0)


def stack(i, items, col, tint):
    n = len(items); y0 = top + 0.3; hh = (bot - y0 - 0.08 * (n - 1)) / n
    for k, it in enumerate(items):
        b = box(s3, xs[i], y0 + k * (hh + 0.08), colw, hh, fill=tint, line=col, lw=1)
        fill_tf(b.text_frame, [{"t": it[0], "bold": True, "size": 9.5, "color": col, "align": PP_ALIGN.CENTER, "after": 1},
                               {"t": it[1], "size": 8.5, "align": PP_ALIGN.CENTER, "after": 0}], anchor=MSO_ANCHOR.MIDDLE, margin=0.04)


stack(0, [("HRMS Adapter", "duty, leave, deployment, transfer, training"), ("Mobile App", "check-in, reaction test, heart rate"),
          ("Wearable", "optional, Health Connect")], MUTED, RGBColor(0xF4, 0xF5, 0xF7))
stack(1, [("Privacy Gateway", "consent check · pseudonymize · minimize · purpose tag · encrypt")], NAVY, TINT)
stack(2, [("Personal baseline", "+ cohort cold-start"), ("Anomaly + change-point", "Isolation Forest"),
          ("4 risk models", "XGBoost + calibration"), ("Explain", "SHAP + counterfactuals")], SAFF, TINT_S)
stack(3, [("Alert engine", "tiered, SLA, escalation"), ("Recommender", "ranked welfare actions"),
          ("Roster optimizer", "OR-Tools, constraints")], TEAL, RGBColor(0xEC, 0xF6, 0xF5))
stack(4, [("Welfare Officer", "pseudonymous cases"), ("Commander", "DP unit aggregates"),
          ("Personnel", "My Trend + Data Mirror")], GREEN, TINT_G)
for i in range(4):
    arrow(s3, xs[i] + colw + 0.015, (top + 0.3 + bot) / 2 - 0.15, w=gap - 0.03)

gate = box(s3, X0, 4.72, 8.1, 0.44, fill=TINT_S, line=SAFF, radius=0.2)
fill_tf(gate.text_frame, [{"t": "**Confidence gate:** low confidence → app requests one optional 90-sec readiness check → evidence fused → confidence 40% → 97%  |  **Loop:** follow-up outcomes feed recovery tracking",
                           "size": 9.5, "align": PP_ALIGN.CENTER, "after": 0}], anchor=MSO_ANCHOR.MIDDLE, margin=0.08)

text(s3, X0, 5.24, 5.0, 0.26, [{"t": "Prototype status: built and tested", "bold": True, "size": 10.5, "color": NAVY}], margin=0)
text(s3, X0 + 4.1, 5.24, 4.0, 0.26, [{"t": "*held-out synthetic personnel", "size": 8.5, "color": MUTED, "align": PP_ALIGN.RIGHT}], margin=0)
built = [("Data", "300 personnel x 180 days synthetic HRMS"), ("AI", "4 risk models, AUROC 0.82-0.97*"),
         ("Explain", "SHAP factors + counterfactual advice"), ("Act", "Tiered alerts, SLA escalation, OR-Tools roster"),
         ("Protect", "DP aggregates, audit chain, purpose lock"), ("Quality", "14 automated tests passing")]
pw = (8.1 - 5 * 0.08) / 6
for i, (h, d) in enumerate(built):
    b = box(s3, X0 + i * (pw + 0.08), 5.52, pw, 1.26, fill=TINT_G, line=RGBColor(0xBF, 0xDD, 0xC4))
    fill_tf(b.text_frame, [{"t": "✓ " + h, "bold": True, "size": 11, "color": GREEN, "after": 2},
                           {"t": d, "size": 9.5, "after": 0}], anchor=MSO_ANCHOR.TOP, margin=0.07)
s3.notes_slide.notes_text_frame.text = (
    "Left: fully open-source, CPU-only, on-prem stack. Right: data flows left to right through the Privacy Gateway before any AI. "
    "Highlight the confidence gate: the system asks for more voluntary evidence instead of over-alerting.")

# ---------------------------------------------------------------- SLIDE 4: feasibility
s4 = S[3]
set_team(s4)
rm(find(s4, "TextBox 8"))
label(s4, 0.4, 1.2, 4.1, "Analysis of the feasibility of the idea", size=14)
feas = [
    ("FaCode", "Technical", "Working prototype runs end to end today. CPU-only open-source ML; AUROC 0.82-0.97 on held-out synthetic personnel.", NAVY),
    ("FaUsers", "Operational", "10-sec daily check-in. Tiered alerts keep load workable (~7% High in a normal unit). Kiosk mode for remote posts.", TEAL),
    ("FaRupeeSign", "Economic", "No licence fees; runs on existing force data centres; one platform reusable across all CAPFs.", SAFF),
    ("FaBalanceScale", "Legal & ethical", "DPDP Act 2023 aligned consent; non-diagnostic; human-in-the-loop; welfare-only purpose lock.", GREEN),
]
y = 1.62
for ic, h, d, col in feas:
    box(s4, 0.4, y, 4.1, 1.2, fill=TINT, line=LINE)
    icon(s4, 0.52, y + 0.12, 0.56, ic, fill=col)
    text(s4, 1.2, y + 0.06, 3.25, 1.1, [{"t": h + "  ✔", "bold": True, "size": 12, "color": col, "after": 2},
                                        {"t": d, "size": 11, "after": 0}], margin=0.02)
    y += 1.3

X4 = 4.75; cwl, cwr = 3.2, 4.38
label(s4, X4, 1.2, cwl + 0.4, "Potential challenges and risks", size=13, color=RED)
label(s4, X4 + cwl + 0.5, 1.2, cwr, "Strategies for overcoming these challenges", size=13, color=GREEN)
risks = [
    ("False alarms & stigma", "Confidence gate + multi-signal agreement + human review; supportive, non-diagnostic language"),
    ("No real labelled data", "Synthetic generator with hidden strain state; governed pilot with force psychologists (DIPR)"),
    ("Low adoption & trust", "10-sec check-in, personal benefit first, My Data Mirror, anonymous support option"),
    ("Data breach or misuse", "On-prem, AES-256, RBAC + ABAC, purpose lock, differential privacy, tamper-evident audit"),
    ("Remote posts / no phone", "Offline sync, shared kiosk with PIN, HRMS-only mode flagged as low confidence"),
    ("Welfare vs mission needs", "Optimizer keeps coverage, skill & rest rules; advisory only, commander approves"),
]
y = 1.62; rh = 0.66
for r, st in risks:
    b = box(s4, X4, y, cwl, rh - 0.08, fill=RGBColor(0xFB, 0xEE, 0xEC), line=RGBColor(0xE8, 0xC4, 0xBE))
    fill_tf(b.text_frame, [{"t": r, "bold": True, "size": 11, "color": RED, "after": 0}], anchor=MSO_ANCHOR.MIDDLE, margin=0.14)
    icon(s4, X4 + cwl - 0.5, y + 0.14, 0.44, "FaExclamationTriangle", fill=RED)
    arrow(s4, X4 + cwl + 0.1, y + (rh - 0.08) / 2 - 0.14, w=0.3, h=0.28, col=MUTED)
    b2 = box(s4, X4 + cwl + 0.5, y, cwr, rh - 0.08, fill=TINT_G, line=RGBColor(0xBF, 0xDD, 0xC4))
    fill_tf(b2.text_frame, [{"t": st, "size": 10, "after": 0}], anchor=MSO_ANCHOR.MIDDLE, margin=0.12)
    y += rh + 0.06
nd = box(s4, X4, 6.05, 12.93 - X4, 0.78, fill=NAVY, radius=0.12)
fill_tf(nd.text_frame, [
    {"t": "What SAHARA will never do", "bold": True, "size": 11, "color": WHITE, "after": 2},
    {"t": "✗ Read chats, calls or social media   ✗ GPS / camera surveillance   ✗ Face or emotion recognition   "
          "✗ Diagnosis or disciplinary use   ✗ Public cloud or LLM for personnel data", "size": 9.5,
     "color": RGBColor(0xD9, 0xE3, 0xEF), "after": 0}], anchor=MSO_ANCHOR.MIDDLE, margin=0.15)
s4.notes_slide.notes_text_frame.text = (
    "Be honest: accuracy is shown on simulated data only; real validation needs a governed pilot. "
    "Each risk on the left has a concrete, already-designed mitigation on the right.")

# ---------------------------------------------------------------- SLIDE 5: impact
s5 = S[4]
set_team(s5)
rm(find(s5, "TextBox 8"))
label(s5, 0.4, 1.2, 9.0, "Potential impact on the target audience", size=14)
aud = [
    ("FaUserShield", "Personnel & families", "Early, confidential support; fairer duty spread; a voice without fear of stigma", NAVY),
    ("FaUserMd", "Welfare Officers & Counsellors", "Prioritized, explainable queue; guided actions; follow-ups tracked automatically", TEAL),
    ("FaUserTie", "Unit Commanders", "Fatigue hotspots 14 days ahead; balanced rosters; readiness without personal data", SAFF),
    ("FaBuilding", "MoHA & Force HQ", "Evidence-based welfare planning, counsellor allocation and retention insight", GREEN),
]
aw = (9.1 - 3 * 0.15) / 4
for i, (ic, h, d, col) in enumerate(aud):
    x = 0.4 + i * (aw + 0.15)
    box(s5, x, 1.6, aw, 2.1, fill=TINT, line=LINE)
    icon(s5, x + 0.15, 1.72, 0.62, ic, fill=col)
    text(s5, x + 0.8, 1.72, aw - 0.88, 0.62, [{"t": h, "bold": True, "size": 10.5, "color": col}], anchor=MSO_ANCHOR.MIDDLE, margin=0)
    text(s5, x + 0.12, 2.42, aw - 0.24, 1.25, [{"t": d, "size": 10.5}], margin=0)

label(s5, 0.4, 3.85, 9.0, "Benefits of the solution (social, economic, environmental, etc.)", size=14)
ben = [
    ("FaHandsHelping", "Social", "Preventive care, not reactive", ["Fewer stress-related incidents", "Stigma-free help-seeking", "Family welfare needs resolved faster"], NAVY),
    ("FaRupeeSign", "Economic", "Protects trained manpower", ["Better retention, lower attrition cost", "Fewer absence & medical days", "Zero licence cost, existing servers"], SAFF),
    ("FaFlag", "Operational", "Stronger force readiness", ["Fatigue-aware, fair rosters", "Balanced night duty & deployments", "Data-driven welfare planning"], TEAL),
    ("FaGlobeAsia", "Environmental & scale", "Green, sovereign, reusable", ["Paperless welfare workflows", "No new hardware; on-prem Indian stack", "Scales to State Police, NDRF/SDRF, Armed Forces"], GREEN),
]
for i, (ic, h, tag, pts, col) in enumerate(ben):
    x = 0.4 + i * (aw + 0.15)
    box(s5, x, 4.25, aw, 2.58, fill=WHITE, line=col, lw=1.25)
    icon(s5, x + 0.15, 4.37, 0.5, ic, fill=col)
    text(s5, x + 0.72, 4.35, aw - 0.8, 0.56, [{"t": h, "bold": True, "size": 11, "color": col, "after": 0},
                                               {"t": tag, "italic": True, "size": 9.5, "color": MUTED, "after": 0}], anchor=MSO_ANCHOR.MIDDLE, margin=0)
    text(s5, x + 0.12, 5.02, aw - 0.24, 1.75, [B(p, size=10.5, after=6, bucol="%02X%02X%02X" % (col[0], col[1], col[2])) for p in pts], margin=0.02)
PH = 5.0
PW = PH * 777 / 1609
px = 9.75 + (12.93 - 9.75 - PW) / 2
s5.shapes.add_picture("phone_crop.png", Inches(px), Inches(1.25), Inches(PW), Inches(PH))
cap5 = box(s5, 9.75, 6.3, 12.93 - 9.75, 0.53, fill=NAVY, radius=0.15)
fill_tf(cap5.text_frame, [{"t": "Live prototype: personnel app in Hindi", "bold": True, "size": 9.5, "color": WHITE, "after": 0, "align": PP_ALIGN.CENTER},
                          {"t": "10-sec check-in · reaction test · help button", "size": 8.5, "color": RGBColor(0xD9, 0xE3, 0xEF),
                           "after": 0, "align": PP_ALIGN.CENTER}], anchor=MSO_ANCHOR.MIDDLE, margin=0.06)
s5.notes_slide.notes_text_frame.text = (
    "Impact story: from reactive to preventive welfare, aimed at the 730 suicides and 55,555 exits reported for 2020-24. "
    "Do not quote real-world reduction percentages; the prototype runs on synthetic data.")

# ---------------------------------------------------------------- SLIDE 6: research
s6 = S[5]
set_team(s6)
rm(find(s6, "TextBox 8"))
label(s6, 0.4, 1.2, 8.0, "Details / Links of the reference and research work", size=14)
refs = [
    ("Problem evidence (Government of India)", [
        "MHA, Rajya Sabha Unstarred Q.1036 (4 Dec 2024): Attrition and suicides amongst CAPFs. mha.gov.in",
    ]),
    ("Validated instruments", [
        "Topp CW et al. The WHO-5 Well-Being Index: a systematic review. Psychother Psychosom. 2015;84:167–176.",
        "Kristensen TS et al. The Copenhagen Burnout Inventory. Work & Stress. 2005;19(3):192–207.",
        "Basner M, Mollicone D, Dinges DF. Validity of a brief psychomotor vigilance test (PVT-B). Acta Astronautica. 2011;69:949–959.",
    ]),
    ("Stress signals & occupational context", [
        "Kim HG et al. Stress and heart rate variability: a meta-analysis. Psychiatry Investig. 2018;15(3):235–245.",
        "Coppetti T et al. Accuracy of smartphone apps for heart rate measurement. Eur J Prev Cardiol. 2017;24(12):1287–1293.",
        "Violanti JM et al. Police stressors and health: a state-of-the-art review. Policing. 2017;40(4):642–656.",
    ]),
    ("AI / ML & privacy methods", [
        "Chen T, Guestrin C. XGBoost: a scalable tree boosting system. KDD 2016.  |  Liu FT et al. Isolation Forest. IEEE ICDM 2008.",
        "Lundberg SM, Lee SI. A unified approach to interpreting model predictions (SHAP). NeurIPS 2017.",
        "Wachter S et al. Counterfactual explanations without opening the black box. Harvard J Law & Tech. 2018.",
        "Dwork C, Roth A. The Algorithmic Foundations of Differential Privacy. 2014.",
    ]),
    ("Policy & platforms (India)", [
        "Digital Personal Data Protection Act, 2023: meity.gov.in",
        "Tele-MANAS national tele-mental health helpline 14416: telemanas.mohfw.gov.in",
        "Google OR-Tools CP-SAT solver: developers.google.com/optimization",
    ]),
]
paras = []
n = 1
for grp, items in refs:
    paras.append({"t": grp, "bold": True, "size": 12.5, "color": NAVY, "after": 3})
    for it in items:
        paras.append({"t": f"[{n}] {it}", "size": 10.5, "after": 3})
        n += 1
box(s6, 0.4, 1.6, 7.9, 5.23, fill=TINT, line=LINE)
text(s6, 0.52, 1.68, 7.7, 5.1, paras, margin=0.04)

xr = 8.5; wr = 4.43
card(s6, xr, 1.6, wr, 5.23, "Gap analysis from our research", [
    {"t": "MHA's current measures [1] → what SAHARA adds", "bold": True, "size": 11.5, "color": SAFF, "after": 4},
    B("**Transparent leave policy** → leave-deprivation tracking and alerts", size=11, after=5),
    B("**Regulating duty hours** → fair-roster optimizer with rest rules", size=11, after=5),
    B("**Officer–troop grievance interaction** → Welfare Need module with SLA", size=11, after=5),
    B("**Stress talks, yoga, counselling** → early, targeted referral to who needs it", size=11, after=12),
    {"t": "Typical tech approach → SAHARA", "bold": True, "size": 11.5, "color": SAFF, "after": 4},
    B("**Needs wearables** → works on any phone + existing HRMS data", size=11, after=5),
    B("**One 'stress score'** → 4 calibrated risks, each with confidence", size=11, after=5),
    B("**Chat / face monitoring** → no surveillance; consent per data type", size=11, after=5),
    B("**Alert only** → counterfactual advice + fair roster + recovery loop", size=11, after=5),
    B("**Claimed accuracy** → measured on held-out data, published honestly", size=11, after=5),
], hcol=NAVY, fill=TINT, iconname="FaBalanceScale", size=9.5)
s6.notes_slide.notes_text_frame.text = (
    "Every design decision is grounded in published research. Instruments are free to use. We never claim clinical accuracy.")

prs.save(OUT)
print("saved", OUT)
